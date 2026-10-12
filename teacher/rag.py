"""Course-scoped, persistent retrieval with bounded, retryable document ingestion.

Heavy model packages are loaded only on first use. Public methods are synchronous:
call them from a bounded worker/thread pool, never on the ASGI event loop. Documents
are immutable content-addressed versions; changed content must be explicitly deleted
when replacing an old version. The application supports one ingestion worker process.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
import threading
import unicodedata
import uuid
import zipfile
from pathlib import PurePosixPath
from typing import Any


class RagValidationError(ValueError):
    """An unsupported, malformed, empty or oversized input document/query."""


class RagService:
    def __init__(self, settings: Any, *, client=None, embedder=None, reranker=None):
        from qdrant_client import QdrantClient, models

        self.settings = settings
        self._models = models
        self.client = (
            client
            if client is not None
            else QdrantClient(
                url=settings.qdrant_url,
                api_key=settings.qdrant_api_key,
                timeout=settings.rag_qdrant_timeout,
            )
        )
        # Same-sized embeddings from different models are not interchangeable.
        signature = f'{settings.rag_embedding_model}@{settings.rag_embedding_revision or "default"}'
        self.collection = (
            f'{settings.qdrant_collection}_{hashlib.sha256(signature.encode()).hexdigest()[:12]}'
        )
        self._embedder = embedder
        self._reranker = reranker
        self._model_lock = threading.RLock()
        self._write_lock = threading.RLock()
        self._collection_lock = threading.Lock()
        self._dimension: int | None = None

    @staticmethod
    def _course(course_id: str) -> str:
        if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]{0,99}', course_id):
            raise RagValidationError('Invalid course identifier')
        return course_id

    def _filter(self, course_id: str, document_id: str | None = None, *, ready=True):
        m = self._models
        conditions = [m.FieldCondition(key='course_id', match=m.MatchValue(value=self._course(course_id)))]
        if document_id is not None:
            try:
                document_id = str(uuid.UUID(document_id))
            except (ValueError, AttributeError) as exc:
                raise RagValidationError('Invalid document identifier') from exc
            conditions.append(m.FieldCondition(key='document_id', match=m.MatchValue(value=document_id)))
        if ready is not None:
            conditions.append(m.FieldCondition(key='ready', match=m.MatchValue(value=ready)))
        return m.Filter(must=conditions)

    def _load_embedder(self):
        with self._model_lock:
            if self._embedder is None:
                from sentence_transformers import SentenceTransformer

                self._embedder = SentenceTransformer(
                    self.settings.rag_embedding_model,
                    revision=self.settings.rag_embedding_revision,
                    device=self.settings.rag_device,
                    trust_remote_code=False,
                )
            return self._embedder

    def _encode(self, texts: list[str], *, query=False) -> list[list[float]]:
        with self._model_lock:
            model = self._load_embedder()
            method = getattr(model, 'encode_query' if query else 'encode_document', model.encode)
            encoded = method(
                texts,
                batch_size=self.settings.rag_batch_size,
                normalize_embeddings=True,
                show_progress_bar=False,
            )
            vectors = [[float(v) for v in row] for row in encoded]
        if len(vectors) != len(texts) or not vectors or not vectors[0]:
            raise RuntimeError('Embedding provider returned an invalid batch')
        dimension = len(vectors[0])
        if any(
            len(row) != dimension or not all(math.isfinite(v) for v in row) or not any(row) for row in vectors
        ):
            raise RuntimeError('Embedding provider returned invalid vectors')
        self._ensure_collection(dimension)
        return vectors

    def _ensure_collection(self, dimension: int):
        with self._collection_lock:
            if self._dimension is not None:
                if self._dimension != dimension:
                    raise RuntimeError('Embedding dimension changed during service lifetime')
                return
            m = self._models
            if not self.client.collection_exists(self.collection):
                try:
                    self.client.create_collection(
                        collection_name=self.collection,
                        vectors_config=m.VectorParams(size=dimension, distance=m.Distance.COSINE),
                    )
                except Exception:
                    # A simultaneous first query may create the same collection.
                    if not self.client.collection_exists(self.collection):
                        raise
            config = self.client.get_collection(self.collection).config.params.vectors
            if isinstance(config, dict) or config.size != dimension or config.distance != m.Distance.COSINE:
                raise RuntimeError('Qdrant collection has incompatible vector configuration')
            for field in ('course_id', 'document_id'):
                self.client.create_payload_index(
                    collection_name=self.collection,
                    field_name=field,
                    field_schema=m.PayloadSchemaType.KEYWORD,
                    wait=True,
                )
            self.client.create_payload_index(
                collection_name=self.collection,
                field_name='ready',
                field_schema=m.PayloadSchemaType.BOOL,
                wait=True,
            )
            self.client.create_payload_index(
                collection_name=self.collection,
                field_name='chunk_index',
                field_schema=m.PayloadSchemaType.INTEGER,
                wait=True,
            )
            self._dimension = dimension

    def _extract(self, filename: str, data: bytes) -> tuple[str, list[tuple[int | None, str]]]:
        if not data or len(data) > self.settings.max_upload_bytes:
            raise RagValidationError('Document is empty or exceeds the upload byte limit')
        title = PurePosixPath(filename.replace('\\', '/')).name
        title = ''.join(
            c for c in unicodedata.normalize('NFKC', title) if not unicodedata.category(c).startswith('C')
        )
        if not title or len(title) > 200:
            raise RagValidationError('Filename must contain between 1 and 200 characters')
        suffix = PurePosixPath(title).suffix.lower()
        max_chars = self.settings.rag_max_extracted_chars
        parts: list[tuple[int | None, str]] = []
        char_count = 0

        def add(page, text):
            nonlocal char_count
            char_count += len(text)
            if char_count > max_chars:
                raise RagValidationError('Document exceeds the extracted text limit')
            text = text.replace('\x00', '').strip()
            if text:
                parts.append((page, text))

        try:
            if suffix in {'.txt', '.md'}:
                add(None, data.decode('utf-8-sig'))
            elif suffix == '.pdf':
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(data), strict=True)
                if reader.is_encrypted:
                    raise RagValidationError('Encrypted PDFs are unsupported')
                if len(reader.pages) > self.settings.rag_max_pdf_pages:
                    raise RagValidationError('PDF exceeds the page limit')
                for index, page in enumerate(reader.pages, 1):
                    add(index, page.extract_text() or '')
            elif suffix == '.docx':
                from docx import Document

                # Bound ZIP expansion before python-docx loads its XML/media parts.
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    infos = archive.infolist()
                    if len(infos) > 2000 or sum(i.file_size for i in infos) > max_chars * 40:
                        raise RagValidationError('DOCX exceeds the expanded archive limit')
                    if 'word/document.xml' not in archive.namelist():
                        raise RagValidationError('Invalid DOCX document')
                document = Document(io.BytesIO(data))
                for paragraph in document.paragraphs:
                    add(None, paragraph.text)
                for table in document.tables:
                    for row in table.rows:
                        add(None, ' | '.join(cell.text for cell in row.cells))
                # Preserve paragraph structure, avoiding one tiny vector per paragraph.
                parts = [(None, '\n\n'.join(text for _, text in parts))] if parts else []
            else:
                raise RagValidationError('Supported formats: UTF-8 TXT, Markdown, PDF and DOCX')
        except RagValidationError:
            raise
        except Exception as exc:
            raise RagValidationError('Document could not be parsed; check its format and encoding') from exc
        if not parts:
            raise RagValidationError('No extractable text found; scanned PDFs require OCR before upload')
        return title, parts

    def _chunks(self, parts: list[tuple[int | None, str]]) -> list[dict]:
        # Fast tokenizers mutate padding/truncation state during model.encode().
        # Serialize tokenizer operations with inference to avoid concurrent borrows.
        with self._model_lock:
            return self._chunks_locked(parts)

    def _chunks_locked(self, parts: list[tuple[int | None, str]]) -> list[dict]:
        model = self._load_embedder()
        tokenizer = model.tokenizer
        limit = min(self.settings.rag_chunk_tokens, int(model.max_seq_length) - 2)
        overlap = self.settings.rag_chunk_overlap
        if limit <= overlap:
            raise RagValidationError('Chunk overlap exceeds the embedding model token window')
        chunks = []
        for page, text in parts:
            tokens = tokenizer.encode(text, add_special_tokens=False, truncation=False)
            start = 0
            while start < len(tokens):
                end = min(start + limit, len(tokens))
                content = tokenizer.decode(tokens[start:end], skip_special_tokens=True).strip()
                # Prefer paragraph/sentence boundaries where at least half a window fits.
                if end < len(tokens):
                    boundaries = list(re.finditer(r'\n\s*\n|[.!?。！？](?:\s|$)', content))
                    if boundaries:
                        candidate = content[: boundaries[-1].end()].strip()
                        candidate_size = len(tokenizer.encode(candidate, add_special_tokens=False))
                        if limit // 2 <= candidate_size < end - start:
                            content, end = candidate, start + candidate_size
                if content:
                    chunks.append({'page': page, 'text': content})
                if len(chunks) > self.settings.rag_max_chunks:
                    raise RagValidationError('Document exceeds the chunk limit')
                if end >= len(tokens):
                    break
                start = max(start + 1, end - overlap)
        if not chunks:
            raise RagValidationError('No indexable text found')
        return chunks

    def ingest(self, course_id: str, filename: str, data: bytes) -> dict:
        self._course(course_id)
        title, parts = self._extract(filename, data)
        digest = hashlib.sha256(data).hexdigest()
        version = f'v1:{self.settings.rag_chunk_tokens}:{self.settings.rag_chunk_overlap}'
        document_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f'{course_id}:{title}:{digest}:{version}'))
        with self._write_lock:
            if self.client.collection_exists(self.collection):
                existing, _ = self.client.scroll(
                    collection_name=self.collection,
                    scroll_filter=self._filter(course_id, document_id),
                    limit=1,
                    with_payload=True,
                    with_vectors=False,
                )
                if existing:
                    payload = existing[0].payload
                    return {
                        'document_id': document_id,
                        'title': title,
                        'chunks': payload['chunk_count'],
                        'sha256': digest,
                        'reused': True,
                    }
            chunks = self._chunks(parts)
            # All embeddings succeed before storage is touched; reject partial model output.
            vectors = self._encode([chunk['text'] for chunk in chunks])
            m = self._models
            points = [
                m.PointStruct(
                    id=str(uuid.uuid5(uuid.UUID(document_id), str(index))),
                    vector=vector,
                    payload={
                        **chunk,
                        'course_id': course_id,
                        'document_id': document_id,
                        'title': title,
                        'sha256': digest,
                        'chunk_index': index,
                        'chunk_count': len(chunks),
                        'ready': False,
                    },
                )
                for index, (chunk, vector) in enumerate(zip(chunks, vectors))
            ]
            for offset in range(0, len(points), self.settings.rag_batch_size):
                self.client.upsert(
                    collection_name=self.collection,
                    points=points[offset : offset + self.settings.rag_batch_size],
                    wait=True,
                )
            # Failed/interrupted ingestion remains invisible; retries overwrite stable IDs.
            self.client.set_payload(
                collection_name=self.collection,
                payload={'ready': True},
                points=self._filter(course_id, document_id, ready=None),
                wait=True,
            )
            return {
                'document_id': document_id,
                'title': title,
                'chunks': len(chunks),
                'sha256': digest,
                'reused': False,
            }

    def search(self, course_id: str, query: str) -> list[dict]:
        return self.search_many([(course_id, query)])[0]

    def search_many(self, requests: list[tuple[str, str]]) -> list[list[dict]]:
        """Batch model/transport work while preserving a filter and budget per query.

        Results correspond positionally to requests. A backend protocol/scope error
        fails the whole batch; no partial or differently scoped data is returned.
        """
        if not requests:
            return []
        if len(requests) > 128:
            raise RagValidationError('Retrieval batch exceeds 128 requests')
        filters, queries = [], []
        for course_id, query in requests:
            filters.append(self._filter(course_id))
            query = query.strip()
            if not query or len(query) > 12000:
                raise RagValidationError('Query must contain between 1 and 12000 characters')
            queries.append(query)
        if not self.client.collection_exists(self.collection):
            return [[] for _ in requests]
        vectors = self._encode(queries, query=True)
        responses = self.client.query_batch_points(
            collection_name=self.collection,
            requests=[
                self._models.QueryRequest(
                    query=vector,
                    filter=query_filter,
                    limit=self.settings.rag_retrieve_k,
                    with_payload=True,
                    score_threshold=self.settings.rag_min_score,
                )
                for vector, query_filter in zip(vectors, filters)
            ],
        )
        if len(responses) != len(requests):
            raise RuntimeError('Vector backend returned an invalid batch size')
        groups = []
        for (course_id, _), response in zip(requests, responses):
            candidates = []
            for hit in response.points:
                p = hit.payload or {}
                # Defense in depth: never trust positional batching without scope checks.
                if p.get('course_id') != course_id or p.get('ready') is not True:
                    raise RuntimeError('Vector backend returned an out-of-scope document')
                if not math.isfinite(float(hit.score)):
                    raise RuntimeError('Vector backend returned an invalid score')
                candidates.append(
                    {
                        'id': str(hit.id),
                        'document_id': p['document_id'],
                        'title': p['title'],
                        'page': p.get('page'),
                        'text': p['text'],
                        'score': float(hit.score),
                        'retrieval_score': float(hit.score),
                    }
                )
            groups.append(candidates)
        pairs = [
            (query, candidate['text'])
            for query, candidates in zip(queries, groups)
            for candidate in candidates
        ]
        if pairs and self.settings.rag_rerank_enabled:
            with self._model_lock:
                if self._reranker is None:
                    from sentence_transformers import CrossEncoder

                    self._reranker = CrossEncoder(
                        self.settings.rag_reranker_model,
                        revision=self.settings.rag_reranker_revision,
                        device=self.settings.rag_device,
                        trust_remote_code=False,
                        max_length=512,
                    )
                scores = self._reranker.predict(
                    pairs,
                    batch_size=self.settings.rag_batch_size,
                    show_progress_bar=False,
                )
            if len(scores) != len(pairs) or any(not math.isfinite(float(score)) for score in scores):
                raise RuntimeError('Reranker returned invalid scores')
            for candidate, score in zip((c for group in groups for c in group), scores):
                candidate['score'] = float(score)
            for candidates in groups:
                candidates.sort(key=lambda item: item['score'], reverse=True)
        with self._model_lock:
            return [self._budget_context(candidates) for candidates in groups]

    def _budget_context(self, candidates: list[dict]) -> list[dict]:
        tokenizer = self._load_embedder().tokenizer
        remaining = self.settings.rag_context_tokens
        results, seen = [], set()
        for candidate in candidates:
            if len(results) >= self.settings.rag_top_k:
                break
            if candidate['text'] in seen:
                continue
            seen.add(candidate['text'])
            # Reserve source-label/metadata overhead in addition to excerpt tokens.
            overhead = 48 + len(tokenizer.encode(candidate['title'], add_special_tokens=False))
            allowance = remaining - overhead
            if allowance < 16:
                break
            tokens = tokenizer.encode(candidate['text'], add_special_tokens=False)
            if len(tokens) > allowance:
                candidate['text'] = tokenizer.decode(tokens[:allowance], skip_special_tokens=True).strip()
            remaining -= overhead + min(len(tokens), allowance)
            results.append(candidate)
        return results

    def documents(self, course_id: str) -> list[dict]:
        query_filter = self._filter(course_id)
        query_filter.must.append(
            self._models.FieldCondition(
                key='chunk_index',
                match=self._models.MatchValue(value=0),
            )
        )
        if not self.client.collection_exists(self.collection):
            return []
        documents = {}
        offset = None
        while True:
            points, offset = self.client.scroll(
                collection_name=self.collection,
                scroll_filter=query_filter,
                limit=256,
                offset=offset,
                with_vectors=False,
                with_payload=['course_id', 'document_id', 'title', 'sha256', 'chunk_count'],
            )
            for point in points:
                p = point.payload
                if p['course_id'] != course_id:
                    raise RuntimeError('Vector backend returned an out-of-scope document')
                documents[p['document_id']] = {
                    'document_id': p['document_id'],
                    'title': p['title'],
                    'sha256': p['sha256'],
                    'chunks': p['chunk_count'],
                }
            if offset is None:
                break
        return sorted(documents.values(), key=lambda doc: (doc['title'], doc['document_id']))

    def delete(self, course_id: str, document_id: str):
        query_filter = self._filter(course_id, document_id, ready=None)
        with self._write_lock:
            if self.client.collection_exists(self.collection):
                self.client.delete(
                    collection_name=self.collection,
                    points_selector=self._models.FilterSelector(filter=query_filter),
                    wait=True,
                )

    def health(self) -> bool:
        try:
            self.client.get_collections()
            return True
        except Exception:
            return False

    def close(self):
        self.client.close()
