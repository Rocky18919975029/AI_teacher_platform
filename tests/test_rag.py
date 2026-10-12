"""Exercise real local Qdrant; models are deterministic and never downloaded."""

import io
import math
import os
import re
from types import SimpleNamespace

import pytest
from qdrant_client import QdrantClient, models

from teacher.rag import RagService, RagValidationError

pytestmark = pytest.mark.filterwarnings('ignore:Payload indexes have no effect:UserWarning')


class Tokenizer:
    def __init__(self):
        self.words = {}
        self.lookup = {}

    def encode(self, text, **kwargs):
        ids = []
        for word in re.findall(r'\w+|[^\w\s]', text):
            if word not in self.words:
                self.words[word] = len(self.words)
                self.lookup[self.words[word]] = word
            ids.append(self.words[word])
        return ids

    def decode(self, tokens, **kwargs):
        return ' '.join(self.lookup[token] for token in tokens)


class Embedder:
    max_seq_length = 512

    def __init__(self):
        self.tokenizer = Tokenizer()
        self.calls = 0

    def encode(self, texts, **kwargs):
        self.calls += 1
        vectors = []
        for text in texts:
            text = text.lower()
            vector = [float(text.count(word)) for word in ('apple', 'banana', 'graph')]
            vector.append(0.05)
            norm = math.sqrt(sum(value * value for value in vector))
            vectors.append([value / norm for value in vector])
        return vectors


@pytest.fixture
def settings():
    return SimpleNamespace(
        qdrant_url='http://unused.invalid',
        qdrant_api_key=None,
        qdrant_collection='test_course',
        rag_qdrant_timeout=2,
        rag_embedding_model='test/embedding',
        rag_embedding_revision='fixed',
        rag_reranker_model='test/reranker',
        rag_reranker_revision='fixed',
        rag_rerank_enabled=False,
        rag_device='cpu',
        rag_chunk_tokens=32,
        rag_chunk_overlap=5,
        rag_retrieve_k=20,
        rag_top_k=5,
        rag_context_tokens=1800,
        rag_min_score=0.25,
        rag_max_extracted_chars=500000,
        rag_max_chunks=2500,
        rag_max_pdf_pages=300,
        rag_batch_size=2,
        max_upload_bytes=10 * 1024 * 1024,
    )


@pytest.fixture
def rag(settings):
    service = RagService(settings, client=QdrantClient(':memory:'), embedder=Embedder())
    yield service
    service.close()


def test_course_isolation_applies_to_search_list_delete_and_ids(rag):
    own = rag.ingest('biology', 'course.md', b'Apple trees are plants.')
    other = rag.ingest('physics', 'course.md', b'Apple trees are plants.')
    assert own['document_id'] != other['document_id']
    assert [d['document_id'] for d in rag.documents('biology')] == [own['document_id']]
    assert {hit['document_id'] for hit in rag.search('biology', 'apple')} == {own['document_id']}
    rag.delete('biology', other['document_id'])
    assert rag.documents('physics')[0]['document_id'] == other['document_id']
    rag.delete('biology', own['document_id'])
    assert rag.documents('biology') == []
    assert rag.search('biology', 'apple') == []
    assert rag.documents('physics')


def test_identical_upload_is_idempotent_and_content_change_versions(rag):
    first = rag.ingest('course', 'notes.txt', b'Apple facts.')
    calls = rag._embedder.calls
    second = rag.ingest('course', 'notes.txt', b'Apple facts.')
    assert second == {**first, 'reused': True}
    assert rag._embedder.calls == calls
    assert len(rag.documents('course')) == 1
    changed = rag.ingest('course', 'notes.txt', b'Banana facts.')
    assert changed['document_id'] != first['document_id']
    assert len(rag.documents('course')) == 2


def test_partial_batch_failure_is_invisible_and_retry_recovers(rag, monkeypatch):
    rag.settings.rag_batch_size = 1
    original = rag.client.upsert
    calls = 0

    def fail_second(**kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise ConnectionError('simulated vector service outage')
        return original(**kwargs)

    monkeypatch.setattr(rag.client, 'upsert', fail_second)
    content = ('apple ' * 100).encode()
    with pytest.raises(ConnectionError):
        rag.ingest('course', 'notes.txt', content)
    assert rag.documents('course') == []
    assert rag.search('course', 'apple') == []
    monkeypatch.setattr(rag.client, 'upsert', original)
    result = rag.ingest('course', 'notes.txt', content)
    assert result['chunks'] > 1
    assert len(rag.documents('course')) == 1
    assert rag.client.count(rag.collection, exact=True).count == result['chunks']


def test_token_chunks_overlap_and_carry_page_provenance(rag):
    chunks = rag._chunks([(7, ' '.join(f'word{i}' for i in range(90)))])
    assert len(chunks) == 4
    assert all(c['page'] == 7 and len(c['text'].split()) <= 32 for c in chunks)
    assert chunks[0]['text'].split()[-5:] == chunks[1]['text'].split()[:5]
    rag.settings.rag_max_chunks = 2
    with pytest.raises(RagValidationError, match='chunk limit'):
        rag._chunks([(7, ' '.join(f'word{i}' for i in range(90)))])


def test_context_budget_trims_excerpts_without_losing_source_metadata(rag):
    result = rag.ingest('course', 'notes.txt', ('apple ' * 100).encode())
    rag.settings.rag_context_tokens = 80
    hits = rag.search('course', 'apple')
    tokenizer = rag._embedder.tokenizer
    assert hits and hits[0]['document_id'] == result['document_id']
    used = sum(48 + len(tokenizer.encode(h['title'])) + len(tokenizer.encode(h['text'])) for h in hits)
    assert used <= 80
    assert all(h['id'] and h['title'] == 'notes.txt' and 'page' in h for h in hits)


def test_reranking_reorders_candidates_and_preserves_dense_score(rag):
    class Reranker:
        def predict(self, pairs, **kwargs):
            return [10 if 'banana' in text.lower() else -10 for _, text in pairs]

    rag.settings.rag_rerank_enabled = True
    rag.settings.rag_min_score = -1
    rag._reranker = Reranker()
    rag.ingest('course', 'a.txt', b'Apple facts.')
    banana = rag.ingest('course', 'b.txt', b'Banana facts.')
    hits = rag.search('course', 'apple')
    assert hits[0]['document_id'] == banana['document_id']
    assert hits[0]['score'] == 10
    assert hits[0]['retrieval_score'] < hits[1]['retrieval_score']


def test_backend_scope_violation_fails_closed(rag, monkeypatch):
    rag.ingest('course', 'notes.txt', b'Apple facts.')
    hit = SimpleNamespace(id='x', score=1, payload={'course_id': 'other', 'ready': True})
    monkeypatch.setattr(rag.client, 'query_points', lambda **kwargs: SimpleNamespace(points=[hit]))
    with pytest.raises(RuntimeError, match='out-of-scope'):
        rag.search('course', 'apple')


@pytest.mark.parametrize(
    'filename,content',
    [
        ('notes.exe', b'hello'),
        ('notes.txt', b'\xff\xfe'),
        ('notes.pdf', b'not a PDF'),
        ('notes.docx', b'not a ZIP'),
        ('notes.txt', b''),
        ('notes.md', b'  \n'),
    ],
)
def test_invalid_input_never_creates_collection(rag, filename, content):
    with pytest.raises(RagValidationError):
        rag.ingest('course', filename, content)
    assert not rag.client.collection_exists(rag.collection)


def test_upload_and_extracted_text_limits(rag):
    rag.settings.max_upload_bytes = 5
    with pytest.raises(RagValidationError, match='byte limit'):
        rag.ingest('course', 'notes.txt', b'abcdef')
    rag.settings.max_upload_bytes = 100
    rag.settings.rag_max_extracted_chars = 4
    with pytest.raises(RagValidationError, match='text limit'):
        rag.ingest('course', 'notes.txt', b'abcdef')


def test_filename_sanitization_and_utf8_bom(rag):
    result = rag.ingest('course', '../private/notes.txt', b'\xef\xbb\xbfApple facts')
    assert result['title'] == 'notes.txt'
    assert rag.search('course', 'apple')[0]['text'] == 'Apple facts'


def test_docx_extracts_paragraphs_and_tables(rag):
    from docx import Document

    document = Document()
    document.add_paragraph('Apple grows on trees.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Banana'
    table.cell(0, 1).text = 'Fruit'
    buffer = io.BytesIO()
    document.save(buffer)
    title, parts = rag._extract('course.docx', buffer.getvalue())
    assert title == 'course.docx'
    assert 'Apple' in parts[0][1] and 'Banana | Fruit' in parts[0][1]


def test_scanned_pdf_and_page_limit_are_explicit_errors(rag):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.add_blank_page(width=300, height=300)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(RagValidationError, match='OCR'):
        rag.ingest('course', 'scan.pdf', buffer.getvalue())
    rag.settings.rag_max_pdf_pages = 1
    with pytest.raises(RagValidationError, match='page limit'):
        rag.ingest('course', 'scan.pdf', buffer.getvalue())


def test_incompatible_collection_rejected_before_writing(rag):
    rag.client.create_collection(
        rag.collection,
        vectors_config=models.VectorParams(size=99, distance=models.Distance.COSINE),
    )
    with pytest.raises(RuntimeError, match='incompatible vector'):
        rag.ingest('course', 'notes.txt', b'Apple')
    assert rag.client.count(rag.collection).count == 0


def test_model_revision_changes_collection_identity(settings):
    client = QdrantClient(':memory:')
    first = RagService(settings, client=client, embedder=Embedder()).collection
    settings.rag_embedding_revision = 'next'
    second = RagService(settings, client=client, embedder=Embedder()).collection
    assert first != second
    client.close()


def test_invalid_scope_query_and_document_ids_rejected(rag):
    with pytest.raises(RagValidationError):
        rag.search('../course', 'apple')
    with pytest.raises(RagValidationError):
        rag.search('course', ' ')
    with pytest.raises(RagValidationError):
        rag.search('course', 'a' * 12001)
    with pytest.raises(RagValidationError):
        rag.delete('course', 'not-a-uuid')


def test_health_tracks_qdrant_failure(rag, monkeypatch):
    assert rag.health() is True

    def fail():
        raise ConnectionError('down')

    monkeypatch.setattr(rag.client, 'get_collections', fail)
    assert rag.health() is False


def test_invalid_embeddings_never_write_partial_records(rag, monkeypatch):
    monkeypatch.setattr(rag._embedder, 'encode', lambda texts, **kwargs: [[float('nan'), 0, 0, 0]])
    with pytest.raises(RuntimeError, match='invalid vectors'):
        rag.ingest('course', 'notes.txt', b'Apple')
    assert not rag.client.collection_exists(rag.collection)


def test_low_relevance_returns_no_context(rag):
    rag.ingest('course', 'notes.txt', b'Apple facts')
    assert rag.search('course', 'graph') == []


def test_reranker_error_propagates_instead_of_silent_quality_change(rag):
    class BrokenReranker:
        def predict(self, *args, **kwargs):
            raise RuntimeError('model unavailable')

    rag.ingest('course', 'notes.txt', b'Apple facts')
    rag.settings.rag_rerank_enabled = True
    rag._reranker = BrokenReranker()
    with pytest.raises(RuntimeError, match='model unavailable'):
        rag.search('course', 'apple')


def test_index_survives_reopen(settings, tmp_path):
    embedder = Embedder()
    first = RagService(settings, client=QdrantClient(path=str(tmp_path)), embedder=embedder)
    document = first.ingest('course', 'notes.txt', b'Apple facts')
    first.close()
    restored = RagService(settings, client=QdrantClient(path=str(tmp_path)), embedder=embedder)
    try:
        assert restored.documents('course')[0]['document_id'] == document['document_id']
        assert restored.search('course', 'apple')[0]['document_id'] == document['document_id']
        assert restored.ingest('course', 'notes.txt', b'Apple facts')['reused'] is True
    finally:
        restored.close()


@pytest.mark.integration
@pytest.mark.skipif(os.getenv('RUN_MODEL_SMOKE') != '1', reason='Set RUN_MODEL_SMOKE=1 to load public models')
def test_real_models_retrieve_rerank_isolate_and_persist(tmp_path):
    """Needs cached public models or access to Hugging Face; no vLLM/GPU required."""
    from teacher.config import Settings

    settings = Settings(
        api_keys=[
            {
                'key': 'smoke-test-key-' * 4,
                'subject': 'smoke',
                'role': 'admin',
                'courses': ['course-a', 'course-b'],
            }
        ],
        rag_rerank_enabled=True,
    )
    service = RagService(settings, client=QdrantClient(path=str(tmp_path)))
    try:
        document = service.ingest(
            'course-a',
            'photosynthesis.md',
            (
                b'Photosynthesis converts sunlight, water, and carbon dioxide into glucose and oxygen '
                b'in chloroplasts. Chlorophyll absorbs light energy. Plants use glucose to grow.'
            ),
        )
        service.ingest(
            'course-a', 'algebra.md', b'A quadratic equation is ax squared plus bx plus c equals zero.'
        )
        private = service.ingest('course-b', 'private.md', b'Photosynthesis course private lecture notes.')
        hits = service.search('course-a', 'How do plants convert sunlight into chemical energy?')
        assert hits and hits[0]['document_id'] == document['document_id']
        assert all(h['document_id'] != private['document_id'] for h in hits)
        assert hits[0]['score'] != hits[0]['retrieval_score']
        embedder, reranker = service._embedder, service._reranker
    finally:
        service.close()
    restored = RagService(
        settings,
        client=QdrantClient(path=str(tmp_path)),
        embedder=embedder,
        reranker=reranker,
    )
    try:
        assert len(restored.documents('course-a')) == 2
        assert (
            restored.search('course-a', 'What does chlorophyll absorb?')[0]['document_id']
            == document['document_id']
        )
    finally:
        restored.close()
